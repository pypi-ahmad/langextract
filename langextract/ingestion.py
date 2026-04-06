# Copyright 2025 Google LLC.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Digital input normalization layer for LangExtract.

This module handles machine-readable inputs and normalizes them into
``NormalizedInput`` objects before they reach the extraction pipeline.

Supported digital ingestion paths in this step:
- plain text strings
- local text files
- digital PDFs with extractable text
- CSV/XLSX/DataFrame/record inputs serialized into structured table text
- safe text and CSV URL retrieval
- image files and scanned PDFs when an OCR engine is provided

OCR-required inputs remain opt-in and require an explicit OCR engine.
"""

from __future__ import annotations

from collections.abc import Iterable, Iterator, Mapping
import csv
import dataclasses
import enum
import functools
import importlib
import io
import itertools
import os
import pathlib
import tempfile
from typing import Any, TYPE_CHECKING

import pandas as pd
import requests

from langextract import ingestion_backends
from langextract import io as lx_io
from langextract.core import data
from langextract.core import exceptions
from langextract.ingestion_backends import FileCategory
from langextract.ingestion_backends import IngestionBackend
from langextract.ingestion_backends import ParserBackendOptions
from langextract.ingestion_backends import register_handler

if TYPE_CHECKING:
  from langextract.ocr import OcrEngine

__all__ = [
    "IngestionError",
    "PdfTextExtractionError",
    "UnsupportedIngestionError",
    "UrlFetchError",
    "InputSourceType",
    "NormalizedContentKind",
    "NormalizedInput",
    "normalize",
    "normalize_input",
]


class NormalizedContentKind(str, enum.Enum):
  """Kinds of payload the extraction flow can consume."""

  TEXT = "text"
  DOCUMENTS = "documents"


class InputSourceType(str, enum.Enum):
  """Detected source category for a normalized input."""

  RAW_TEXT = "raw_text"
  DOCUMENT = "document"
  DOCUMENTS = "documents"
  LOCAL_PATH = "local_path"
  URL = "url"
  BYTES = "bytes"
  PDF = "pdf"
  CSV = "csv"
  XLSX = "xlsx"
  TABLE = "table"
  IMAGE = "image"


@dataclasses.dataclass(slots=True, frozen=True)
class NormalizedInput:
  """Common internal representation for user-provided inputs."""

  source_type: InputSourceType
  content_kind: NormalizedContentKind
  text: str | None = None
  documents: Iterable[data.Document] | None = None
  source_name: str | None = None
  metadata: dict[str, Any] = dataclasses.field(default_factory=dict)

  def __post_init__(self) -> None:
    has_text = self.text is not None
    has_documents = self.documents is not None

    if has_text == has_documents:
      raise ValueError(
          "NormalizedInput must contain exactly one of 'text' or 'documents'."
      )

    if self.content_kind == NormalizedContentKind.TEXT and not has_text:
      raise ValueError("TEXT inputs must provide text.")
    if (
        self.content_kind == NormalizedContentKind.DOCUMENTS
        and not has_documents
    ):
      raise ValueError("DOCUMENTS inputs must provide documents.")

  @classmethod
  def from_text(
      cls,
      text: str,
      *,
      source_type: InputSourceType,
      source_name: str | None = None,
      metadata: dict[str, Any] | None = None,
  ) -> NormalizedInput:
    return cls(
        source_type=source_type,
        content_kind=NormalizedContentKind.TEXT,
        text=text,
        source_name=source_name,
        metadata=dict(metadata or {}),
    )

  @classmethod
  def from_documents(
      cls,
      documents: Iterable[data.Document],
      *,
      source_type: InputSourceType,
      source_name: str | None = None,
      metadata: dict[str, Any] | None = None,
  ) -> NormalizedInput:
    return cls(
        source_type=source_type,
        content_kind=NormalizedContentKind.DOCUMENTS,
        documents=documents,
        source_name=source_name,
        metadata=dict(metadata or {}),
    )

  @property
  def is_text(self) -> bool:
    return self.content_kind == NormalizedContentKind.TEXT

  @property
  def is_documents(self) -> bool:
    return self.content_kind == NormalizedContentKind.DOCUMENTS

  def as_extract_input(self) -> str | Iterable[data.Document]:
    if self.is_text:
      return self.text or ""
    return self.documents or []


class IngestionError(exceptions.LangExtractError):
  """Raised when input normalization fails."""


class UnsupportedIngestionError(IngestionError):
  """Raised when the input type is recognized but not supported yet."""


class UrlFetchError(IngestionError):
  """Raised when URL retrieval fails or returns unusable content."""


class PdfTextExtractionError(IngestionError):
  """Raised when a PDF cannot provide machine-readable text."""


@dataclasses.dataclass(frozen=True, slots=True)
class _PdfBackendSelection:
  readable_backend: str | None = None
  scanned_backend: str | None = None
  table_backend: str | None = None


@dataclasses.dataclass(frozen=True, slots=True)
class _PdfParseResult:
  text: str | None
  metadata: dict[str, Any]


_PDF_BACKEND_CATEGORIES: tuple[FileCategory, ...] = (
    FileCategory.READABLE_PDF,
    FileCategory.SCANNED_PDF,
    FileCategory.TABLE_PDF,
)


def _resolve_pdf_backend_selection(
    *,
    backend: str | None,
    parser_backends: ParserBackendOptions | None,
) -> _PdfBackendSelection:
  readable_backend = None
  scanned_backend = None
  table_backend = None

  if parser_backends is not None:
    readable_backend = parser_backends.readable_pdf_backend
    scanned_backend = parser_backends.scanned_pdf_backend
    table_backend = parser_backends.table_pdf_backend

  if backend is None:
    return _PdfBackendSelection(
        readable_backend=readable_backend,
        scanned_backend=scanned_backend,
        table_backend=table_backend,
    )

  if backend == "auto":
    return _PdfBackendSelection(
        readable_backend="auto",
        scanned_backend="auto",
        table_backend=table_backend,
    )

  matches = [
      category
      for category in _PDF_BACKEND_CATEGORIES
      if ingestion_backends.get_backend(category, backend) is not None
  ]
  if not matches:
    raise ValueError(
        f"Unknown PDF parser backend {backend!r}. Configure PDF backends "
        "with ParserBackendOptions(readable_pdf_backend=..., "
        "scanned_pdf_backend=..., table_pdf_backend=...)."
    )
  if len(matches) > 1:
    raise ValueError(
        f"Ambiguous PDF parser backend {backend!r}; it matches categories "
        f"{[category.value for category in matches]}. Configure PDF "
        "backends with ParserBackendOptions(readable_pdf_backend=..., "
        "scanned_pdf_backend=..., table_pdf_backend=...) to choose "
        "explicitly."
    )

  category = matches[0]
  if category == FileCategory.READABLE_PDF:
    readable_backend = backend
  elif category == FileCategory.SCANNED_PDF:
    scanned_backend = backend
  else:
    table_backend = backend

  return _PdfBackendSelection(
      readable_backend=readable_backend,
      scanned_backend=scanned_backend,
      table_backend=table_backend,
  )


_EMPTY = object()

_TEXT_EXTENSIONS: frozenset[str] = frozenset({
    ".txt",
    ".text",
    ".md",
    ".rst",
    ".log",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
})

_HTML_EXTENSIONS: frozenset[str] = frozenset({
    ".html",
    ".htm",
})

_IMAGE_EXTENSIONS: frozenset[str] = frozenset({
    ".png",
    ".jpg",
    ".jpeg",
    ".tiff",
    ".tif",
    ".bmp",
    ".webp",
    ".gif",
})

_TEXT_CONTENT_TYPES: tuple[str, ...] = (
    "text/",
    "application/json",
    "application/xml",
    "application/yaml",
    "application/x-yaml",
)

_CSV_CONTENT_TYPES: frozenset[str] = frozenset({
    "text/csv",
    "application/csv",
    "application/vnd.ms-excel",
})

_REMOTE_BINARY_CONTENT_TYPES: frozenset[str] = frozenset({
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "image/png",
    "image/jpeg",
    "image/tiff",
    "image/webp",
    "image/gif",
    "image/bmp",
})


# ---------------------------------------------------------------------------
# Extension → category mapping (used by table-driven dispatch in _from_path)
# ---------------------------------------------------------------------------

_EXTENSION_CATEGORIES: dict[str, FileCategory] = {}
for _ext in _TEXT_EXTENSIONS:
  _EXTENSION_CATEGORIES[_ext] = FileCategory.TXT
for _ext in _HTML_EXTENSIONS:
  _EXTENSION_CATEGORIES[_ext] = FileCategory.HTML
_EXTENSION_CATEGORIES[".csv"] = FileCategory.CSV
for _ext in (".xlsx", ".xls"):
  _EXTENSION_CATEGORIES[_ext] = FileCategory.EXCEL
_EXTENSION_CATEGORIES[".docx"] = FileCategory.DOCX
_EXTENSION_CATEGORIES[".doc"] = FileCategory.DOC
_EXTENSION_CATEGORIES[".pptx"] = FileCategory.PPTX
_EXTENSION_CATEGORIES[".ppt"] = FileCategory.PPT
for _ext in _IMAGE_EXTENSIONS:
  _EXTENSION_CATEGORIES[_ext] = FileCategory.IMAGE
# Note: .pdf handled separately by _from_pdf (multi-category runtime detection)


def normalize_input(
    input_data: Any,
    *,
    text_column: str = "text",
    id_column: str | None = None,
    additional_context_column: str | None = None,
    fetch_urls: bool = True,
    ocr_engine: OcrEngine | None = None,
    backend: str | None = None,
    parser_backends: ParserBackendOptions | None = None,
) -> NormalizedInput:
  """Normalize supported inputs into the common internal model.

  Args:
    input_data: The raw user input to normalize.
    text_column: Column name for text in tabular inputs.
    id_column: Column name for document IDs in tabular inputs.
    additional_context_column: Column used as per-row context.
    fetch_urls: Whether to fetch URL strings.
    ocr_engine: Optional OCR engine for image/scanned-PDF inputs.
    backend: Optional ingestion backend name to use instead of the
      default. Must be a backend registered in
      ``langextract.ingestion_backends`` for the detected file category.
      When ``None`` (the default), the category's default backend is used.
    parser_backends: Optional category-specific backend selections. These are
      applied after the input type is detected and preserve existing behavior
      when omitted.
  """
  if isinstance(input_data, NormalizedInput):
    return input_data

  if isinstance(input_data, str):
    if fetch_urls and lx_io.is_url(input_data):
      return _from_url(
          input_data,
          text_column=text_column,
          id_column=id_column,
          additional_context_column=additional_context_column,
          backend=backend,
          parser_backends=parser_backends,
      )
    return NormalizedInput.from_text(
        input_data,
        source_type=InputSourceType.RAW_TEXT,
    )

  if isinstance(input_data, data.Document):
    return NormalizedInput.from_documents(
        [input_data],
        source_type=InputSourceType.DOCUMENT,
    )

  if isinstance(input_data, os.PathLike):
    return _from_path(
        pathlib.Path(input_data),
        text_column=text_column,
        id_column=id_column,
        additional_context_column=additional_context_column,
        ocr_engine=ocr_engine,
        backend=backend,
        parser_backends=parser_backends,
    )

  if isinstance(input_data, pd.DataFrame):
    return _from_dataframe(
        input_data,
        text_column=text_column,
        id_column=id_column,
        additional_context_column=additional_context_column,
    )

  if isinstance(input_data, (list, tuple)):
    if not input_data:
      raise IngestionError(
          "Empty iterables are not supported. Pass at least one Document or "
          "one mapping record."
      )
    if all(isinstance(item, data.Document) for item in input_data):
      return NormalizedInput.from_documents(
          input_data,
          source_type=InputSourceType.DOCUMENTS,
      )
    if all(isinstance(item, Mapping) for item in input_data):
      return _from_records(
          input_data,
          text_column=text_column,
          id_column=id_column,
          additional_context_column=additional_context_column,
      )
    raise IngestionError(
        "Mixed iterables are not supported. Pass either all Document objects "
        "or all mapping records."
    )

  if isinstance(input_data, Mapping):
    return _from_records(
        [input_data],
        text_column=text_column,
        id_column=id_column,
        additional_context_column=additional_context_column,
    )

  if isinstance(input_data, (bytes, bytearray, memoryview)):
    return _from_bytes(
        bytes(input_data),
        ocr_engine=ocr_engine,
        backend=backend,
        parser_backends=parser_backends,
    )

  iterable_result = _from_iterable(
      input_data,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )
  if iterable_result is not None:
    return iterable_result

  raise _unsupported_input_error(input_data)


def normalize(
    input_data: Any,
    *,
    text_column: str = "text",
    id_column: str | None = None,
    additional_context_column: str | None = None,
    fetch_urls: bool = True,
    ocr_engine: OcrEngine | None = None,
    backend: str | None = None,
    parser_backends: ParserBackendOptions | None = None,
) -> str | Iterable[data.Document]:
  """Compatibility wrapper returning the legacy extract payload shapes."""
  normalized = normalize_input(
      input_data,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
      fetch_urls=fetch_urls,
      ocr_engine=ocr_engine,
      backend=backend,
      parser_backends=parser_backends,
  )
  return normalized.as_extract_input()


def _from_path(
    path: pathlib.Path,
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    ocr_engine: OcrEngine | None = None,
    backend: str | None = None,
    parser_backends: ParserBackendOptions | None = None,
) -> NormalizedInput:
  if not path.exists():
    raise FileNotFoundError(f"File not found: {path}")

  suffix = path.suffix.lower()

  # PDF has multi-category runtime detection (readable vs scanned vs table).
  if suffix == ".pdf":
    return _from_pdf(
        path,
        ocr_engine=ocr_engine,
        backend=backend,
        parser_backends=parser_backends,
    )

  category = _EXTENSION_CATEGORIES.get(suffix)
  if category is None:
    raise UnsupportedIngestionError(f"Unsupported file type: {suffix} ({path})")

  explicit_backend_selected = backend is not None or (
      parser_backends is not None
      and parser_backends.backend_for_category(category) is not None
  )

  if (
      category == FileCategory.IMAGE
      and ocr_engine is not None
      and not explicit_backend_selected
  ):
    return _ocr_image_file(path, ocr_engine=ocr_engine)

  # Try registered handler dispatch.
  result = ingestion_backends.resolve_handler(
      category,
      backend=backend,
      parser_backends=parser_backends,
  )
  if result is not None:
    _, handler = result
    return handler(
        path=path,
        text_column=text_column,
        id_column=id_column,
        additional_context_column=additional_context_column,
        ocr_engine=ocr_engine,
    )

  # No dispatch handler — validate any explicit backend selection (may raise
  # ValueError / NotImplementedError / ImportError), then apply fallbacks.
  ingestion_backends.resolve_backend(
      category,
      backend=backend,
      parser_backends=parser_backends,
  )

  if category == FileCategory.IMAGE:
    if ocr_engine is not None:
      return _ocr_image_file(path, ocr_engine=ocr_engine)
    raise UnsupportedIngestionError(
        "Image ingestion requires an OCR engine. Pass an ocr_engine "
        "(e.g. OllamaOcrEngine) to normalize_input() to enable it."
    )

  raise UnsupportedIngestionError(f"Unsupported file type: {suffix} ({path})")


def _from_text_file(
    path: pathlib.Path,
    *,
    format_name: str = "text",
    backend_name: str | None = None,
) -> NormalizedInput:
  raw = path.read_bytes()
  text = _decode_text_bytes(raw)
  if text is None:
    raise IngestionError(
        f"Could not decode text file {path} as UTF-8 or UTF-16 text."
    )

  metadata: dict[str, Any] = {
      "path": str(path),
      "suffix": path.suffix.lower(),
      "format": format_name,
  }
  if backend_name is not None:
    metadata["backend"] = backend_name

  return NormalizedInput.from_text(
      text,
      source_type=InputSourceType.LOCAL_PATH,
      source_name=str(path),
      metadata=metadata,
  )


def _from_csv(
    path: pathlib.Path,
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    backend_name: str = "pandas",
) -> NormalizedInput:
  try:
    df = pd.read_csv(path, dtype=str)
  except pd.errors.EmptyDataError as e:
    raise IngestionError(f"CSV file is empty: {path}") from e
  except Exception as e:  # pragma: no cover - defensive wrapper
    raise IngestionError(f"Failed to read CSV file {path}: {e}") from e

  serialized = _serialize_table(
      df,
      table_label="csv",
      source_name=str(path),
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )
  return NormalizedInput.from_text(
      serialized,
      source_type=InputSourceType.CSV,
      source_name=str(path),
      metadata={
          "path": str(path),
          "columns": list(df.columns),
          "row_count": len(df.index),
          "backend": backend_name,
      },
  )


def _from_excel(
    path: pathlib.Path,
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    backend_name: str = "pandas",
) -> NormalizedInput:
  excel_backend = ingestion_backends.get_backend(FileCategory.EXCEL, "pandas")
  try:
    import openpyxl  # noqa: F401 - dependency presence check
  except ImportError as e:
    if excel_backend is not None:
      raise ingestion_backends._build_missing_dependency_error(
          excel_backend,
          missing_import="openpyxl",
          missing_package="openpyxl",
          extras=("office", "xlsx"),
      ) from e
    raise

  try:
    df = pd.read_excel(path, dtype=str, engine="openpyxl")
  except Exception as e:  # pragma: no cover - defensive wrapper
    raise IngestionError(f"Failed to read XLSX file {path}: {e}") from e

  serialized = _serialize_table(
      df,
      table_label="xlsx",
      source_name=str(path),
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )
  return NormalizedInput.from_text(
      serialized,
      source_type=InputSourceType.XLSX,
      source_name=str(path),
      metadata={
          "path": str(path),
          "columns": list(df.columns),
          "row_count": len(df.index),
          "backend": backend_name,
      },
  )


def _import_backend_module(
    category: FileCategory,
    backend_name: str,
) -> tuple[IngestionBackend, Any]:
  backend = ingestion_backends.require_backend(category, backend_name)
  if backend.import_check is None:
    raise IngestionError(
        f"Ingestion backend {backend.name!r} for category "
        f"{backend.category.value!r} does not expose an importable Python "
        "module."
    )

  try:
    return backend, importlib.import_module(backend.import_check)
  except ImportError as e:
    raise ingestion_backends._build_missing_dependency_error(backend) from e


def _clean_extracted_text(text: str) -> str:
  normalized = text.replace("\r\n", "\n").replace("\r", "\n")
  cleaned_lines: list[str] = []
  previous_blank = True

  for line in normalized.split("\n"):
    stripped = line.strip()
    if stripped:
      cleaned_lines.append(stripped)
      previous_blank = False
    elif not previous_blank:
      cleaned_lines.append("")
      previous_blank = True

  return "\n".join(cleaned_lines).strip()


def _normalized_local_backend_input(
    path: pathlib.Path,
    text: str,
    *,
    format_name: str,
    backend_name: str,
) -> NormalizedInput:
  return NormalizedInput.from_text(
      _clean_extracted_text(text),
      source_type=InputSourceType.LOCAL_PATH,
      source_name=str(path),
      metadata={
          "path": str(path),
          "suffix": path.suffix.lower(),
          "format": format_name,
          "backend": backend_name,
      },
  )


def _normalized_url_backend_input(
    url: str,
    text: str,
    *,
    content_type: str,
    backend_name: str,
) -> NormalizedInput:
  return NormalizedInput.from_text(
      _clean_extracted_text(text),
      source_type=InputSourceType.URL,
      source_name=url,
      metadata={
          "url": url,
          "content_type": content_type or "unknown",
          "format": "html",
          "backend": backend_name,
      },
  )


def _build_pdf_parse_result(page_texts: Iterable[str]) -> _PdfParseResult:
  page_sections: list[str] = []
  text_pages: list[int] = []
  empty_pages: list[int] = []
  materialized_pages = list(page_texts)

  for page_number, page_text in enumerate(materialized_pages, start=1):
    stripped = (page_text or "").strip()
    if not stripped:
      empty_pages.append(page_number)
      continue

    text_pages.append(page_number)
    page_sections.append(f"[Page {page_number}]\n{stripped}")

  return _PdfParseResult(
      text="\n\n".join(page_sections) if page_sections else None,
      metadata={
          "page_count": len(materialized_pages),
          "text_pages": text_pages,
          "empty_pages": empty_pages,
      },
  )


def _readable_pdf_with_pymupdf(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
) -> _PdfParseResult:
  pymupdf = _import_pymupdf()
  try:
    if path is not None:
      doc = pymupdf.open(str(path))
    else:
      doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
  except Exception as e:  # pragma: no cover - defensive wrapper
    target = str(path) if path is not None else "PDF bytes"
    raise PdfTextExtractionError(
        f"Failed to open PDF input {target}: {e}"
    ) from e

  try:
    return _build_pdf_parse_result(
        (doc[page_index].get_text() or "") for page_index in range(len(doc))
    )
  finally:
    doc.close()


def _readable_pdf_with_pdfplumber(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
) -> _PdfParseResult:
  _, pdfplumber = _import_backend_module(
      FileCategory.READABLE_PDF, "pdfplumber"
  )
  source = str(path) if path is not None else io.BytesIO(pdf_bytes or b"")

  try:
    with pdfplumber.open(source) as document:
      return _build_pdf_parse_result(
          (page.extract_text() or "") for page in document.pages
      )
  except Exception as e:  # pragma: no cover - defensive wrapper
    target = str(path) if path is not None else "PDF bytes"
    raise PdfTextExtractionError(
        f"Failed to read PDF input {target} with pdfplumber: {e}"
    ) from e


def _readable_pdf_with_pypdf(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
) -> _PdfParseResult:
  _, pypdf = _import_backend_module(FileCategory.READABLE_PDF, "pypdf")
  source = str(path) if path is not None else io.BytesIO(pdf_bytes or b"")

  try:
    reader = pypdf.PdfReader(source)
    return _build_pdf_parse_result(
        (page.extract_text() or "") for page in reader.pages
    )
  except Exception as e:  # pragma: no cover - defensive wrapper
    target = str(path) if path is not None else "PDF bytes"
    raise PdfTextExtractionError(
        f"Failed to read PDF input {target} with pypdf: {e}"
    ) from e


def _from_csv_builtin(
    *,
    path: pathlib.Path | None = None,
    raw: bytes | None = None,
    source_name: str,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
) -> NormalizedInput:
  if raw is None and path is not None:
    raw = path.read_bytes()
  decoded = _decode_text_bytes(raw or b"")
  if decoded is None:
    raise IngestionError(
        f"Could not decode CSV input {source_name} as UTF-8 or UTF-16 text."
    )

  try:
    rows = list(csv.reader(io.StringIO(decoded)))
  except csv.Error as e:
    raise IngestionError(f"Failed to read CSV input {source_name}: {e}") from e

  if not rows:
    raise IngestionError(f"CSV input is empty: {source_name}")

  headers = ["" if header is None else str(header) for header in rows[0]]
  records = [
      [row[index] if index < len(row) else "" for index in range(len(headers))]
      for row in rows[1:]
  ]
  df = pd.DataFrame(records, columns=headers)

  serialized = _serialize_table(
      df,
      table_label="csv",
      source_name=source_name,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )
  metadata: dict[str, Any] = {
      "columns": list(df.columns),
      "row_count": len(df.index),
      "backend": "builtin_csv",
  }
  if path is not None:
    metadata["path"] = str(path)
  if lx_io.is_url(source_name):
    metadata["url"] = source_name

  return NormalizedInput.from_text(
      serialized,
      source_type=InputSourceType.CSV,
      source_name=source_name,
      metadata=metadata,
  )


def _from_excel_openpyxl(
    path: pathlib.Path,
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
) -> NormalizedInput:
  _, openpyxl = _import_backend_module(FileCategory.EXCEL, "openpyxl")

  try:
    workbook = openpyxl.load_workbook(
        str(path),
        data_only=True,
        read_only=True,
    )
  except Exception as e:  # pragma: no cover - defensive wrapper
    raise IngestionError(f"Failed to read XLSX file {path}: {e}") from e

  try:
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
  finally:
    workbook.close()

  if not rows:
    raise IngestionError(f"XLSX file is empty: {path}")

  headers = ["" if value is None else str(value) for value in rows[0]]
  records = [
      [
          row[index] if index < len(headers) else None
          for index in range(len(headers))
      ]
      for row in rows[1:]
  ]
  df = pd.DataFrame(records, columns=headers)

  serialized = _serialize_table(
      df,
      table_label="xlsx",
      source_name=str(path),
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )
  return NormalizedInput.from_text(
      serialized,
      source_type=InputSourceType.XLSX,
      source_name=str(path),
      metadata={
          "path": str(path),
          "columns": list(df.columns),
          "row_count": len(df.index),
          "backend": "openpyxl",
      },
  )


def _extract_docx_with_python_docx(path: pathlib.Path) -> str:
  _, docx = _import_backend_module(FileCategory.DOCX, "python-docx")

  try:
    document = docx.Document(str(path))
  except Exception as e:  # pragma: no cover - defensive wrapper
    raise IngestionError(f"Failed to read DOCX file {path}: {e}") from e

  sections: list[str] = []
  sections.extend(paragraph.text for paragraph in document.paragraphs)
  for table in document.tables:
    for row in table.rows:
      cell_text = [cell.text.strip() for cell in row.cells]
      if any(cell_text):
        sections.append(" | ".join(cell_text))

  return "\n\n".join(section for section in sections if section.strip())


def _extract_docx_with_docx2txt(path: pathlib.Path) -> str:
  _, docx2txt = _import_backend_module(FileCategory.DOCX, "docx2txt")

  try:
    return docx2txt.process(str(path)) or ""
  except Exception as e:  # pragma: no cover - defensive wrapper
    raise IngestionError(f"Failed to read DOCX file {path}: {e}") from e


def _extract_pptx_with_python_pptx(path: pathlib.Path) -> str:
  _, pptx = _import_backend_module(FileCategory.PPTX, "python-pptx")

  try:
    presentation = pptx.Presentation(str(path))
  except Exception as e:  # pragma: no cover - defensive wrapper
    raise IngestionError(f"Failed to read PPTX file {path}: {e}") from e

  slide_sections: list[str] = []
  for slide_number, slide in enumerate(presentation.slides, start=1):
    text_chunks = [
        str(shape.text)
        for shape in slide.shapes
        if hasattr(shape, "text") and str(shape.text).strip()
    ]
    if text_chunks:
      slide_sections.append(
          f"[Slide {slide_number}]\n" + "\n".join(text_chunks)
      )

  return "\n\n".join(slide_sections)


def _extract_html_with_trafilatura(
    html_text: str,
    *,
    category: FileCategory,
) -> str:
  _, trafilatura = _import_backend_module(category, "trafilatura")
  extracted = trafilatura.extract(
      html_text,
      include_comments=False,
      include_tables=True,
      output_format="txt",
  )
  if extracted is None:
    raise IngestionError(
        "Trafilatura could not extract text from the HTML input."
    )
  return extracted


def _extract_html_with_beautifulsoup(
    html_text: str,
    *,
    category: FileCategory,
) -> str:
  _, bs4 = _import_backend_module(category, "beautifulsoup")
  soup = bs4.BeautifulSoup(html_text, "html.parser")
  for tag in soup(["script", "style", "noscript"]):
    tag.decompose()
  return soup.get_text(separator="\n")


def _decode_html_bytes(raw: bytes, *, source_name: str) -> str:
  html_text = _decode_text_bytes(raw)
  if html_text is None:
    raise IngestionError(
        f"Could not decode HTML input {source_name} as UTF-8 or UTF-16 text."
    )
  return html_text


def _import_pillow(backend: IngestionBackend):
  try:
    from PIL import Image  # pylint: disable=import-outside-toplevel
  except ImportError as e:
    raise ingestion_backends._build_missing_dependency_error(
        backend,
        missing_import="PIL",
        missing_package="pillow",
        extras=("ocr",),
    ) from e
  return Image


def _extract_text_with_tesseract(
    image_data: bytes,
    *,
    category: FileCategory,
) -> str:
  backend, pytesseract = _import_backend_module(category, "tesseract")
  image_lib = _import_pillow(backend)

  try:
    with image_lib.open(io.BytesIO(image_data)) as image:
      return pytesseract.image_to_string(image)
  except Exception as e:  # pragma: no cover - environment dependent
    tesseract_not_found = getattr(pytesseract, "TesseractNotFoundError", None)
    if tesseract_not_found is not None and isinstance(e, tesseract_not_found):
      raise IngestionError(
          "Selected ingestion backend 'tesseract' requires the system "
          "'tesseract' executable to be installed and available on PATH."
      ) from e
    raise IngestionError(f"Tesseract OCR failed: {e}") from e


@functools.lru_cache(maxsize=1)
def _build_paddleocr_instance() -> Any:
  paddleocr = importlib.import_module("paddleocr")
  try:
    return paddleocr.PaddleOCR(use_angle_cls=True, lang="en")
  except TypeError:  # pragma: no cover - older PaddleOCR signatures
    return paddleocr.PaddleOCR(lang="en")


def _extract_paddleocr_text(result: Any) -> str:
  lines: list[str] = []

  def collect(node: Any) -> None:
    if isinstance(node, (list, tuple)):
      if (
          len(node) >= 2
          and isinstance(node[1], (list, tuple))
          and len(node[1]) >= 1
          and isinstance(node[1][0], str)
      ):
        lines.append(node[1][0])
        return
      for item in node:
        collect(item)

  collect(result)
  return "\n".join(lines)


def _extract_text_with_paddleocr_path(
    path: pathlib.Path,
    *,
    category: FileCategory,
) -> str:
  _import_backend_module(category, "paddleocr")
  try:
    return _extract_paddleocr_text(
        _build_paddleocr_instance().ocr(str(path), cls=True)
    )
  except Exception as e:  # pragma: no cover - environment dependent
    raise IngestionError(f"PaddleOCR failed: {e}") from e


def _extract_text_with_paddleocr_bytes(
    image_data: bytes,
    *,
    category: FileCategory,
) -> str:
  _import_backend_module(category, "paddleocr")

  with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
    handle.write(image_data)
    handle.flush()
    temp_path = pathlib.Path(handle.name)

  try:
    return _extract_text_with_paddleocr_path(temp_path, category=category)
  finally:
    temp_path.unlink(missing_ok=True)


def _render_pdf_page_images(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    selected_backend: IngestionBackend | None = None,
) -> list[tuple[int, bytes]]:
  try:
    pymupdf = _import_pymupdf()
  except ImportError as e:
    if selected_backend is not None:
      raise ingestion_backends._build_missing_dependency_error(
          selected_backend,
          missing_import="pymupdf",
          missing_package="pymupdf",
          extras=("pdf",),
      ) from e
    raise

  if path is not None:
    doc = pymupdf.open(str(path))
  else:
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")

  try:
    return [
        (page_index + 1, doc[page_index].get_pixmap().tobytes("png"))
        for page_index in range(len(doc))
    ]
  finally:
    doc.close()


def _ocr_pdf_with_backend(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    source_name: str,
    backend_name: str,
) -> tuple[str, dict[str, Any]]:
  backend = ingestion_backends.require_backend(
      FileCategory.SCANNED_PDF, backend_name
  )
  page_images = _render_pdf_page_images(
      path=path,
      pdf_bytes=pdf_bytes,
      selected_backend=backend,
  )
  page_sections: list[str] = []
  ocr_pages: list[int] = []

  for page_number, image_data in page_images:
    if backend_name == "tesseract":
      extracted = _extract_text_with_tesseract(
          image_data,
          category=FileCategory.SCANNED_PDF,
      )
    else:
      extracted = _extract_text_with_paddleocr_bytes(
          image_data,
          category=FileCategory.SCANNED_PDF,
      )
    text = _clean_extracted_text(extracted)
    if text:
      page_sections.append(f"[Page {page_number}]\n{text}")
      ocr_pages.append(page_number)

  if not page_sections:
    raise PdfTextExtractionError(
        f"{source_name} produced no text with the {backend_name!r} OCR backend."
    )

  return "\n\n".join(page_sections), {
      "page_count": len(page_images),
      "ocr_pages": ocr_pages,
      "format": "ocr",
      "backend": backend_name,
  }


def _from_readable_pdf_backend(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    source_name: str,
    backend_name: str,
) -> NormalizedInput:
  if backend_name == "pymupdf":
    parsed = _readable_pdf_with_pymupdf(path=path, pdf_bytes=pdf_bytes)
  elif backend_name == "pdfplumber":
    parsed = _readable_pdf_with_pdfplumber(path=path, pdf_bytes=pdf_bytes)
  elif backend_name == "pypdf":
    parsed = _readable_pdf_with_pypdf(path=path, pdf_bytes=pdf_bytes)
  else:  # pragma: no cover - resolved by registry validation
    raise UnsupportedIngestionError(
        f"Unsupported readable PDF backend: {backend_name}"
    )

  if parsed.text is None:
    raise PdfTextExtractionError(
        f"{source_name} contains no digitally extractable text."
    )

  metadata = dict(parsed.metadata)
  metadata.update({"format": "digital", "backend": backend_name})
  if path is not None:
    metadata["path"] = str(path)
  else:
    metadata["origin"] = "bytes"

  return NormalizedInput.from_text(
      parsed.text,
      source_type=InputSourceType.PDF,
      source_name=None if path is None else source_name,
      metadata=metadata,
  )


def _from_scanned_pdf_backend(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    source_name: str,
    backend_name: str,
) -> NormalizedInput:
  extracted_text, metadata = _ocr_pdf_with_backend(
      path=path,
      pdf_bytes=pdf_bytes,
      source_name=source_name,
      backend_name=backend_name,
  )
  if path is not None:
    metadata["path"] = str(path)
  else:
    metadata["origin"] = "bytes"

  return NormalizedInput.from_text(
      extracted_text,
      source_type=InputSourceType.PDF,
      source_name=None if path is None else source_name,
      metadata=metadata,
  )


def _from_html_path_backend(
    path: pathlib.Path,
    *,
    backend_name: str,
) -> NormalizedInput:
  html_text = _decode_html_bytes(path.read_bytes(), source_name=str(path))
  if backend_name == "trafilatura":
    extracted = _extract_html_with_trafilatura(
        html_text,
        category=FileCategory.HTML,
    )
  elif backend_name == "beautifulsoup":
    extracted = _extract_html_with_beautifulsoup(
        html_text,
        category=FileCategory.HTML,
    )
  else:  # pragma: no cover - resolved by registry validation
    raise UnsupportedIngestionError(f"Unsupported HTML backend: {backend_name}")

  return _normalized_local_backend_input(
      path,
      extracted,
      format_name="html",
      backend_name=backend_name,
  )


def _from_html_url_backend(
    *,
    url: str,
    raw: bytes,
    content_type: str,
    backend_name: str,
) -> NormalizedInput:
  html_text = _decode_html_bytes(raw, source_name=url)
  if backend_name == "trafilatura":
    extracted = _extract_html_with_trafilatura(
        html_text, category=FileCategory.URL
    )
  elif backend_name == "beautifulsoup":
    extracted = _extract_html_with_beautifulsoup(
        html_text, category=FileCategory.URL
    )
  else:  # pragma: no cover - resolved by registry validation
    raise UnsupportedIngestionError(f"Unsupported URL backend: {backend_name}")

  return _normalized_url_backend_input(
      url,
      extracted,
      content_type=content_type,
      backend_name=backend_name,
  )


def _from_image_path_backend(
    path: pathlib.Path,
    *,
    backend_name: str,
) -> NormalizedInput:
  if backend_name == "tesseract":
    extracted = _extract_text_with_tesseract(
        path.read_bytes(),
        category=FileCategory.IMAGE,
    )
  elif backend_name == "paddleocr":
    extracted = _extract_text_with_paddleocr_path(
        path,
        category=FileCategory.IMAGE,
    )
  else:  # pragma: no cover - resolved by registry validation
    raise UnsupportedIngestionError(
        f"Unsupported image backend: {backend_name}"
    )

  return NormalizedInput.from_text(
      _clean_extracted_text(extracted),
      source_type=InputSourceType.IMAGE,
      source_name=str(path),
      metadata={
          "path": str(path),
          "suffix": path.suffix.lower(),
          "format": "ocr",
          "backend": backend_name,
      },
  )


def _from_image_bytes_backend(
    raw: bytes, *, backend_name: str
) -> NormalizedInput:
  if backend_name == "tesseract":
    extracted = _extract_text_with_tesseract(
        raw,
        category=FileCategory.IMAGE,
    )
  elif backend_name == "paddleocr":
    extracted = _extract_text_with_paddleocr_bytes(
        raw,
        category=FileCategory.IMAGE,
    )
  else:  # pragma: no cover - resolved by registry validation
    raise UnsupportedIngestionError(
        f"Unsupported image backend: {backend_name}"
    )

  return NormalizedInput.from_text(
      _clean_extracted_text(extracted),
      source_type=InputSourceType.IMAGE,
      metadata={
          "origin": "bytes",
          "format": "ocr",
          "backend": backend_name,
      },
  )


# ---------------------------------------------------------------------------
# Handler registrations (table-driven dispatch for _from_path)
# ---------------------------------------------------------------------------


@register_handler(FileCategory.TXT, "builtin")
def _handle_text_file(*, path: pathlib.Path, **_unused: Any) -> NormalizedInput:
  return _from_text_file(path, backend_name="builtin")


@register_handler(FileCategory.HTML, "builtin")
def _handle_html_builtin(
    *, path: pathlib.Path, **_unused: Any
) -> NormalizedInput:
  return _from_text_file(
      path,
      format_name="html",
      backend_name="builtin",
  )


@register_handler(FileCategory.HTML, "trafilatura")
def _handle_html_trafilatura(
    *,
    path: pathlib.Path,
    **_unused: Any,
) -> NormalizedInput:
  return _from_html_path_backend(path, backend_name="trafilatura")


@register_handler(FileCategory.HTML, "beautifulsoup")
def _handle_html_beautifulsoup(
    *,
    path: pathlib.Path,
    **_unused: Any,
) -> NormalizedInput:
  return _from_html_path_backend(path, backend_name="beautifulsoup")


@register_handler(FileCategory.READABLE_PDF, "pymupdf")
def _handle_pdf_pymupdf(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    source_name: str,
    **_unused: Any,
) -> NormalizedInput:
  return _from_readable_pdf_backend(
      path=path,
      pdf_bytes=pdf_bytes,
      source_name=source_name,
      backend_name="pymupdf",
  )


@register_handler(FileCategory.READABLE_PDF, "pdfplumber")
def _handle_pdf_pdfplumber(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    source_name: str,
    **_unused: Any,
) -> NormalizedInput:
  return _from_readable_pdf_backend(
      path=path,
      pdf_bytes=pdf_bytes,
      source_name=source_name,
      backend_name="pdfplumber",
  )


@register_handler(FileCategory.READABLE_PDF, "pypdf")
def _handle_pdf_pypdf(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    source_name: str,
    **_unused: Any,
) -> NormalizedInput:
  return _from_readable_pdf_backend(
      path=path,
      pdf_bytes=pdf_bytes,
      source_name=source_name,
      backend_name="pypdf",
  )


@register_handler(FileCategory.SCANNED_PDF, "tesseract")
def _handle_pdf_tesseract(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    source_name: str,
    **_unused: Any,
) -> NormalizedInput:
  return _from_scanned_pdf_backend(
      path=path,
      pdf_bytes=pdf_bytes,
      source_name=source_name,
      backend_name="tesseract",
  )


@register_handler(FileCategory.SCANNED_PDF, "paddleocr")
def _handle_pdf_paddleocr(
    *,
    path: pathlib.Path | None = None,
    pdf_bytes: bytes | None = None,
    source_name: str,
    **_unused: Any,
) -> NormalizedInput:
  return _from_scanned_pdf_backend(
      path=path,
      pdf_bytes=pdf_bytes,
      source_name=source_name,
      backend_name="paddleocr",
  )


@register_handler(FileCategory.CSV, "pandas")
def _handle_csv(
    *,
    path: pathlib.Path,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    **_unused: Any,
) -> NormalizedInput:
  return _from_csv(
      path,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
      backend_name="pandas",
  )


@register_handler(FileCategory.CSV, "builtin_csv")
def _handle_csv_builtin(
    *,
    path: pathlib.Path,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    **_unused: Any,
) -> NormalizedInput:
  return _from_csv_builtin(
      path=path,
      source_name=str(path),
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )


@register_handler(FileCategory.EXCEL, "pandas")
def _handle_excel(
    *,
    path: pathlib.Path,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    **_unused: Any,
) -> NormalizedInput:
  return _from_excel(
      path,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
      backend_name="pandas",
  )


@register_handler(FileCategory.EXCEL, "openpyxl")
def _handle_excel_openpyxl(
    *,
    path: pathlib.Path,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    **_unused: Any,
) -> NormalizedInput:
  return _from_excel_openpyxl(
      path,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )


@register_handler(FileCategory.DOCX, "python-docx")
def _handle_docx_python_docx(
    *,
    path: pathlib.Path,
    **_unused: Any,
) -> NormalizedInput:
  return _normalized_local_backend_input(
      path,
      _extract_docx_with_python_docx(path),
      format_name="docx",
      backend_name="python-docx",
  )


@register_handler(FileCategory.DOCX, "docx2txt")
def _handle_docx_docx2txt(
    *,
    path: pathlib.Path,
    **_unused: Any,
) -> NormalizedInput:
  return _normalized_local_backend_input(
      path,
      _extract_docx_with_docx2txt(path),
      format_name="docx",
      backend_name="docx2txt",
  )


@register_handler(FileCategory.PPTX, "python-pptx")
def _handle_pptx_python_pptx(
    *,
    path: pathlib.Path,
    **_unused: Any,
) -> NormalizedInput:
  return _normalized_local_backend_input(
      path,
      _extract_pptx_with_python_pptx(path),
      format_name="pptx",
      backend_name="python-pptx",
  )


@register_handler(FileCategory.IMAGE, "tesseract")
def _handle_image_tesseract(
    *,
    path: pathlib.Path,
    **_unused: Any,
) -> NormalizedInput:
  return _from_image_path_backend(path, backend_name="tesseract")


@register_handler(FileCategory.IMAGE, "paddleocr")
def _handle_image_paddleocr(
    *,
    path: pathlib.Path,
    **_unused: Any,
) -> NormalizedInput:
  return _from_image_path_backend(path, backend_name="paddleocr")


@register_handler(FileCategory.URL, "trafilatura")
def _handle_url_trafilatura(
    *,
    url: str,
    raw: bytes,
    content_type: str,
    **_unused: Any,
) -> NormalizedInput:
  return _from_html_url_backend(
      url=url,
      raw=raw,
      content_type=content_type,
      backend_name="trafilatura",
  )


@register_handler(FileCategory.URL, "beautifulsoup")
def _handle_url_beautifulsoup(
    *,
    url: str,
    raw: bytes,
    content_type: str,
    **_unused: Any,
) -> NormalizedInput:
  return _from_html_url_backend(
      url=url,
      raw=raw,
      content_type=content_type,
      backend_name="beautifulsoup",
  )


def _from_dataframe(
    df: pd.DataFrame,
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
) -> NormalizedInput:
  serialized = _serialize_table(
      df,
      table_label="table",
      source_name=None,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )
  return NormalizedInput.from_text(
      serialized,
      source_type=InputSourceType.TABLE,
      metadata={
          "columns": list(df.columns),
          "row_count": len(df.index),
      },
  )


def _from_records(
    records: Iterable[Mapping[str, Any]],
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
) -> NormalizedInput:
  validated_records = list(_validated_mappings(records))
  if not validated_records:
    raise IngestionError(
        "Empty iterables are not supported. Pass at least one mapping record."
    )

  try:
    df = pd.DataFrame.from_records(validated_records)
  except Exception as e:  # pragma: no cover - defensive wrapper
    raise IngestionError(f"Failed to normalize mapping records: {e}") from e

  return _from_dataframe(
      df,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )


def _serialize_table(
    df: pd.DataFrame,
    *,
    table_label: str,
    source_name: str | None,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
) -> str:
  if list(df.columns) == []:
    raise IngestionError("Tabular input has no columns to serialize.")
  if df.empty:
    raise IngestionError("Tabular input contains no rows to serialize.")

  ordered_columns = _ordered_columns(
      list(df.columns),
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )
  view = df.reindex(columns=ordered_columns)

  lines: list[str] = []
  if source_name:
    lines.append(f"Table source: {source_name}")
  else:
    lines.append("Table input")
  lines.append(f"Format: {table_label}")
  lines.append("Columns: " + " | ".join(ordered_columns))
  lines.append("")

  for row_number, row in enumerate(
      view.itertuples(index=False, name=None), start=1
  ):
    row_parts = [
        f"{column}={_quote_table_value(value)}"
        for column, value in zip(ordered_columns, row)
    ]
    lines.append(f"Row {row_number}: " + " | ".join(row_parts))

  return "\n".join(lines)


def _ordered_columns(
    columns: list[str],
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
) -> list[str]:
  prioritized: list[str] = []
  for column_name in (text_column, id_column, additional_context_column):
    if (
        column_name
        and column_name in columns
        and column_name not in prioritized
    ):
      prioritized.append(column_name)

  remaining = [column for column in columns if column not in prioritized]
  return prioritized + remaining


def _quote_table_value(value: Any) -> str:
  if value is None or pd.isna(value):
    return '""'

  text = str(value).replace("\r\n", "\n").replace("\r", "\n")
  text = text.replace("\\", "\\\\").replace('"', '\\"')
  text = text.replace("\n", "\\n")
  return f'"{text}"'


def _import_pymupdf():
  try:
    import pymupdf

    return pymupdf
  except ImportError as e:
    backend = ingestion_backends.get_backend(
        FileCategory.READABLE_PDF, "pymupdf"
    )
    if backend is not None:
      raise ingestion_backends._build_missing_dependency_error(backend) from e
    raise


def _from_pdf(
    path: pathlib.Path,
    ocr_engine: OcrEngine | None = None,
    *,
    backend: str | None = None,
    parser_backends: ParserBackendOptions | None = None,
) -> NormalizedInput:
  pdf_selection = _resolve_pdf_backend_selection(
      backend=backend,
      parser_backends=parser_backends,
  )
  explicit_scanned_backend_selected = pdf_selection.scanned_backend is not None
  if pdf_selection.table_backend is not None:
    ingestion_backends.require_backend(
        FileCategory.TABLE_PDF, pdf_selection.table_backend
    )
  readable_result = None
  readable_resolution_error: ImportError | None = None
  try:
    readable_result = ingestion_backends.resolve_handler(
        FileCategory.READABLE_PDF,
        backend=pdf_selection.readable_backend,
        parser_backends=parser_backends,
    )
  except ImportError as e:
    if (
        pdf_selection.scanned_backend is None
        or pdf_selection.readable_backend not in (None, "auto")
    ):
      raise
    readable_resolution_error = e

  readable_error: PdfTextExtractionError | None = None
  if readable_result is not None:
    _, handler = readable_result
    try:
      return handler(path=path, source_name=str(path))
    except PdfTextExtractionError as e:
      readable_error = e

  if explicit_scanned_backend_selected:
    scanned_result = ingestion_backends.resolve_handler(
        FileCategory.SCANNED_PDF,
        backend=pdf_selection.scanned_backend,
        parser_backends=parser_backends,
    )
    if scanned_result is not None:
      _, handler = scanned_result
      return handler(path=path, source_name=str(path))

  if ocr_engine is not None:
    pymupdf = _import_pymupdf()
    try:
      doc = pymupdf.open(str(path))
    except Exception as e:  # pragma: no cover - defensive wrapper
      raise PdfTextExtractionError(
          f"Failed to open PDF file {path}: {e}"
      ) from e

    try:
      extracted_text, metadata = _ocr_pdf_pages(
          doc,
          source_name=str(path),
          ocr_engine=ocr_engine,
      )
    finally:
      doc.close()

    metadata["path"] = str(path)
    return NormalizedInput.from_text(
        extracted_text,
        source_type=InputSourceType.PDF,
        source_name=str(path),
        metadata=metadata,
    )

  if readable_error is not None:
    scanned_result = ingestion_backends.resolve_handler(
        FileCategory.SCANNED_PDF,
        parser_backends=parser_backends,
    )
    if scanned_result is not None:
      _, handler = scanned_result
      return handler(path=path, source_name=str(path))

  if readable_error is not None:
    if explicit_scanned_backend_selected:
      raise readable_error
    raise PdfTextExtractionError(
        f"{path} contains no digitally extractable text. Pass an ocr_engine "
        "to normalize_input() to enable OCR preprocessing for scanned PDFs."
    ) from readable_error

  if readable_resolution_error is not None:
    raise readable_resolution_error

  raise UnsupportedIngestionError(f"Unsupported PDF ingestion path: {path}")


def _from_pdf_bytes(
    pdf_bytes: bytes,
    ocr_engine: OcrEngine | None = None,
    *,
    backend: str | None = None,
    parser_backends: ParserBackendOptions | None = None,
) -> NormalizedInput:
  pdf_selection = _resolve_pdf_backend_selection(
      backend=backend,
      parser_backends=parser_backends,
  )
  explicit_scanned_backend_selected = pdf_selection.scanned_backend is not None
  if pdf_selection.table_backend is not None:
    ingestion_backends.require_backend(
        FileCategory.TABLE_PDF, pdf_selection.table_backend
    )
  readable_result = None
  readable_resolution_error: ImportError | None = None
  try:
    readable_result = ingestion_backends.resolve_handler(
        FileCategory.READABLE_PDF,
        backend=pdf_selection.readable_backend,
        parser_backends=parser_backends,
    )
  except ImportError as e:
    if (
        pdf_selection.scanned_backend is None
        or pdf_selection.readable_backend not in (None, "auto")
    ):
      raise
    readable_resolution_error = e

  readable_error: PdfTextExtractionError | None = None
  if readable_result is not None:
    _, handler = readable_result
    try:
      return handler(pdf_bytes=pdf_bytes, source_name="PDF bytes")
    except PdfTextExtractionError as e:
      readable_error = e

  if explicit_scanned_backend_selected:
    scanned_result = ingestion_backends.resolve_handler(
        FileCategory.SCANNED_PDF,
        backend=pdf_selection.scanned_backend,
        parser_backends=parser_backends,
    )
    if scanned_result is not None:
      _, handler = scanned_result
      return handler(pdf_bytes=pdf_bytes, source_name="PDF bytes")

  if ocr_engine is not None:
    pymupdf = _import_pymupdf()
    try:
      doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    except Exception as e:  # pragma: no cover - defensive wrapper
      raise PdfTextExtractionError(f"Failed to open PDF bytes: {e}") from e

    try:
      extracted_text, metadata = _ocr_pdf_pages(
          doc,
          source_name="PDF bytes",
          ocr_engine=ocr_engine,
      )
    finally:
      doc.close()

    metadata["origin"] = "bytes"
    return NormalizedInput.from_text(
        extracted_text,
        source_type=InputSourceType.PDF,
        metadata=metadata,
    )

  if readable_error is not None:
    scanned_result = ingestion_backends.resolve_handler(
        FileCategory.SCANNED_PDF,
        parser_backends=parser_backends,
    )
    if scanned_result is not None:
      _, handler = scanned_result
      return handler(pdf_bytes=pdf_bytes, source_name="PDF bytes")

  if readable_error is not None:
    raise PdfTextExtractionError(
        "PDF bytes contain no digitally extractable text. Pass an ocr_engine "
        "to normalize_input() to enable OCR preprocessing for scanned PDFs."
    ) from readable_error

  if readable_resolution_error is not None:
    raise readable_resolution_error

  raise UnsupportedIngestionError("Unsupported PDF byte ingestion path.")


def _extract_pdf_text(
    doc: Any,
    *,
    source_name: str,
    ocr_engine: OcrEngine | None = None,
    scanned_backend: str | None = None,
) -> tuple[str, dict[str, Any]]:
  page_sections: list[str] = []
  text_pages: list[int] = []
  empty_pages: list[int] = []

  for page_index in range(len(doc)):
    page_text = (doc[page_index].get_text() or "").strip()
    page_number = page_index + 1
    if not page_text:
      empty_pages.append(page_number)
      continue

    text_pages.append(page_number)
    page_sections.append(f"[Page {page_number}]\n{page_text}")

  if not page_sections:
    if scanned_backend is not None:
      ingestion_backends.require_backend(
          FileCategory.SCANNED_PDF, scanned_backend
      )
    if ocr_engine is not None:
      return _ocr_pdf_pages(
          doc,
          source_name=source_name,
          ocr_engine=ocr_engine,
      )
    raise PdfTextExtractionError(
        f"{source_name} contains no digitally extractable text. "
        "Pass an ocr_engine to normalize_input() to enable OCR "
        "preprocessing for scanned PDFs."
    )

  return "\n\n".join(page_sections), {
      "page_count": len(doc),
      "text_pages": text_pages,
      "empty_pages": empty_pages,
  }


def _from_url(
    url: str,
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    backend: str | None = None,
    parser_backends: ParserBackendOptions | None = None,
) -> NormalizedInput:
  try:
    response = requests.get(url, timeout=lx_io.DEFAULT_TIMEOUT_SECONDS)
    response.raise_for_status()
  except requests.RequestException as e:
    raise UrlFetchError(f"Failed to fetch URL {url}: {e}") from e

  content_type = response.headers.get("Content-Type", "")
  content_type = content_type.split(";", 1)[0].strip().lower()
  url_path = url.split("?", 1)[0].split("#", 1)[0]
  url_ext = pathlib.PurePosixPath(url_path).suffix.lower()

  if url_ext == ".csv" or content_type in _CSV_CONTENT_TYPES:
    csv_result = ingestion_backends.resolve_handler(
        FileCategory.CSV,
        backend=backend,
        parser_backends=parser_backends,
    )
    if csv_result is not None:
      selected_backend, _ = csv_result
      if selected_backend.name == "builtin_csv":
        return _from_csv_builtin(
            raw=response.content,
            source_name=url,
            text_column=text_column,
            id_column=id_column,
            additional_context_column=additional_context_column,
        )
      return _from_csv_bytes(
          response.content,
          source_name=url,
          text_column=text_column,
          id_column=id_column,
          additional_context_column=additional_context_column,
          backend_name=selected_backend.name,
      )

    return _from_csv_bytes(
        response.content,
        source_name=url,
        text_column=text_column,
        id_column=id_column,
        additional_context_column=additional_context_column,
        backend_name="pandas",
    )

  if _is_text_like_url(url_ext=url_ext, content_type=content_type):
    text_backend: ingestion_backends.IngestionBackend | None = None
    if _is_html_like_url(url_ext=url_ext, content_type=content_type):
      html_result = ingestion_backends.resolve_handler(
          FileCategory.URL,
          backend=backend,
          parser_backends=parser_backends,
      )
      if html_result is not None:
        _, handler = html_result
        return handler(
            url=url,
            raw=response.content,
            content_type=content_type,
        )
    else:
      text_backend = ingestion_backends.resolve_backend(
          FileCategory.TXT,
          backend=backend,
          parser_backends=parser_backends,
      )
    text = _decode_text_bytes(response.content)
    if text is None:
      raise UrlFetchError(
          f"Fetched URL {url} but could not decode the response as text."
      )
    metadata = {
        "url": url,
        "content_type": content_type or "unknown",
    }
    if text_backend is not None:
      metadata["backend"] = text_backend.name
    return NormalizedInput.from_text(
        text,
        source_type=InputSourceType.URL,
        source_name=url,
        metadata=metadata,
    )

  if _is_remote_binary_document(url_ext=url_ext, content_type=content_type):
    # TODO: Add safe remote binary document ingestion for PDF/XLSX once the
    # local digital ingestion path is stable.
    raise UnsupportedIngestionError(
        "Remote PDF, XLSX, and image ingestion is not enabled yet. "
        "Download the file locally and pass pathlib.Path for machine-readable "
        "documents. OCR-backed remote ingestion is deferred to a later step."
    )

  raise UnsupportedIngestionError(
      f"Unsupported URL content for ingestion: {url}. Only text-like URLs "
      "and CSV URLs are currently supported."
  )


def _is_text_like_url(*, url_ext: str, content_type: str) -> bool:
  if url_ext in _TEXT_EXTENSIONS or url_ext in _HTML_EXTENSIONS:
    return True
  return any(content_type.startswith(prefix) for prefix in _TEXT_CONTENT_TYPES)


def _is_html_like_url(*, url_ext: str, content_type: str) -> bool:
  if url_ext in _HTML_EXTENSIONS:
    return True
  return content_type == "text/html"


def _is_remote_binary_document(*, url_ext: str, content_type: str) -> bool:
  if url_ext in _IMAGE_EXTENSIONS:
    return True
  if url_ext in {".pdf", ".xlsx", ".xls"}:
    return True
  return content_type in _REMOTE_BINARY_CONTENT_TYPES


def _from_bytes(
    raw: bytes,
    ocr_engine: OcrEngine | None = None,
    *,
    backend: str | None = None,
    parser_backends: ParserBackendOptions | None = None,
) -> NormalizedInput:
  if raw[:5] == b"%PDF-":
    return _from_pdf_bytes(
        raw,
        ocr_engine=ocr_engine,
        backend=backend,
        parser_backends=parser_backends,
    )

  if _looks_like_image_bytes(raw):
    explicit_image_backend_selected = backend is not None or (
        parser_backends is not None
        and parser_backends.image_backend is not None
    )
    if ocr_engine is not None and not explicit_image_backend_selected:
      return _ocr_image_bytes(raw, ocr_engine=ocr_engine)

    image_result = ingestion_backends.resolve_handler(
        FileCategory.IMAGE,
        backend=backend,
        parser_backends=parser_backends,
    )
    if image_result is not None:
      selected_backend, _ = image_result
      return _from_image_bytes_backend(raw, backend_name=selected_backend.name)
    if ocr_engine is not None:
      return _ocr_image_bytes(raw, ocr_engine=ocr_engine)
    raise UnsupportedIngestionError(
        "Image byte ingestion requires an OCR engine. Pass an ocr_engine "
        "to normalize_input() to enable it."
    )

  decoded_text = _decode_text_bytes(raw)
  if decoded_text is not None:
    text_backend = ingestion_backends.resolve_backend(
        FileCategory.TXT,
        backend=backend,
        parser_backends=parser_backends,
    )
    metadata = {
        "origin": "bytes",
        "format": "text",
    }
    if text_backend is not None:
      metadata["backend"] = text_backend.name
    return NormalizedInput.from_text(
        decoded_text,
        source_type=InputSourceType.BYTES,
        metadata=metadata,
    )

  raise IngestionError(
      "Cannot determine format from raw bytes. Supported byte inputs in this "
      "step are UTF-8/UTF-16 text and digital PDFs. Image and scanned inputs "
      "still require OCR preprocessing."
  )


def _from_iterable(
    input_data: Any,
    *,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
) -> NormalizedInput | None:
  if not isinstance(input_data, Iterable):
    return None

  if isinstance(input_data, (str, bytes, bytearray, memoryview, os.PathLike)):
    return None

  first_item, items = _peek_iterable(input_data)
  if first_item is _EMPTY:
    raise IngestionError(
        "Empty iterables are not supported. Pass at least one Document or "
        "one mapping record."
    )

  if isinstance(first_item, data.Document):
    return NormalizedInput.from_documents(
        _validated_documents(items),
        source_type=InputSourceType.DOCUMENTS,
    )
  if isinstance(first_item, Mapping):
    return _from_records(
        _validated_mappings(items),
        text_column=text_column,
        id_column=id_column,
        additional_context_column=additional_context_column,
    )

  return None


def _peek_iterable(values: Iterable[Any]) -> tuple[object, Iterable[Any]]:
  iterator = iter(values)
  try:
    first_item = next(iterator)
  except StopIteration:
    return _EMPTY, ()
  return first_item, itertools.chain([first_item], iterator)


def _validated_documents(items: Iterable[Any]) -> Iterator[data.Document]:
  for index, item in enumerate(items):
    if not isinstance(item, data.Document):
      raise IngestionError(
          "Mixed iterables are not supported. Expected Document objects, got "
          f"{type(item).__name__} at position {index}."
      )
    yield item


def _validated_mappings(items: Iterable[Any]) -> Iterator[Mapping[str, Any]]:
  for index, item in enumerate(items):
    if not isinstance(item, Mapping):
      raise IngestionError(
          "Mixed iterables are not supported. Expected mapping records, got "
          f"{type(item).__name__} at position {index}."
      )
    yield item


def _from_csv_bytes(
    raw: bytes,
    *,
    source_name: str,
    text_column: str,
    id_column: str | None,
    additional_context_column: str | None,
    backend_name: str = "pandas",
) -> NormalizedInput:
  try:
    df = pd.read_csv(io.BytesIO(raw), dtype=str)
  except pd.errors.EmptyDataError as e:
    raise IngestionError(f"CSV input is empty: {source_name}") from e
  except Exception as e:  # pragma: no cover - defensive wrapper
    raise IngestionError(f"Failed to read CSV input {source_name}: {e}") from e

  serialized = _serialize_table(
      df,
      table_label="csv",
      source_name=source_name,
      text_column=text_column,
      id_column=id_column,
      additional_context_column=additional_context_column,
  )
  metadata: dict[str, Any] = {
      "columns": list(df.columns),
      "row_count": len(df.index),
      "backend": backend_name,
  }
  if lx_io.is_url(source_name):
    metadata["url"] = source_name

  return NormalizedInput.from_text(
      serialized,
      source_type=InputSourceType.CSV,
      source_name=source_name,
      metadata=metadata,
  )


def _looks_like_image_bytes(raw: bytes) -> bool:
  return any([
      raw.startswith(b"\x89PNG\r\n\x1a\n"),
      raw.startswith(b"\xff\xd8"),
      raw.startswith((b"GIF87a", b"GIF89a")),
      raw.startswith((b"II*\x00", b"MM\x00*")),
      raw.startswith(b"BM"),
      raw.startswith(b"RIFF") and raw[8:12] == b"WEBP",
  ])


def _decode_text_bytes(raw: bytes) -> str | None:
  if not raw:
    return ""

  encodings = ["utf-8-sig", "utf-8"]
  if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
    encodings.append("utf-16")

  for encoding in encodings:
    try:
      decoded = raw.decode(encoding)
    except UnicodeDecodeError:
      continue
    if _is_probably_text(decoded):
      return decoded
  return None


def _is_probably_text(text: str) -> bool:
  if not text:
    return True

  printable = sum(ch.isprintable() or ch in "\r\n\t" for ch in text)
  return printable / len(text) >= 0.9


def _ocr_image_file(
    path: pathlib.Path,
    ocr_engine: OcrEngine,
) -> NormalizedInput:
  """Read an image file and run OCR on it."""
  image_data = path.read_bytes()
  result = ocr_engine.run_ocr(image_data)
  return NormalizedInput.from_text(
      result.text,
      source_type=InputSourceType.IMAGE,
      source_name=str(path),
      metadata={
          "path": str(path),
          "suffix": path.suffix.lower(),
          "format": "ocr",
          "ocr_metadata": result.metadata,
      },
  )


def _ocr_image_bytes(
    raw: bytes,
    ocr_engine: OcrEngine,
) -> NormalizedInput:
  """Run OCR on raw image bytes."""
  result = ocr_engine.run_ocr(raw)
  return NormalizedInput.from_text(
      result.text,
      source_type=InputSourceType.IMAGE,
      metadata={
          "origin": "bytes",
          "format": "ocr",
          "ocr_metadata": result.metadata,
      },
  )


def _ocr_pdf_pages(
    doc: Any,
    *,
    source_name: str,
    ocr_engine: OcrEngine,
) -> tuple[str, dict[str, Any]]:
  """Render every PDF page to an image and OCR it."""
  page_sections: list[str] = []
  ocr_pages: list[int] = []

  for page_index in range(len(doc)):
    page_number = page_index + 1
    pixmap = doc[page_index].get_pixmap()
    image_data = pixmap.tobytes("png")
    result = ocr_engine.run_ocr_pdf_page(
        image_data,
        page_number=page_number,
    )
    text = result.text.strip()
    if text:
      page_sections.append(f"[Page {page_number}]\n{text}")
      ocr_pages.append(page_number)

  if not page_sections:
    raise PdfTextExtractionError(
        f"{source_name} produced no text even after OCR."
    )

  return "\n\n".join(page_sections), {
      "page_count": len(doc),
      "ocr_pages": ocr_pages,
      "format": "ocr",
  }


def _unsupported_input_error(input_data: Any) -> IngestionError:
  type_name = type(input_data).__name__
  return IngestionError(
      "Unsupported input type: "
      f"{type_name}. Supported inputs include raw text strings, http(s) URL "
      "strings, pathlib.Path/os.PathLike values, Document objects, iterables "
      "of Document objects, table-like records (mapping or iterable of "
      "mappings), pandas DataFrames, and bytes/bytearray/memoryview payloads. "
      "If you mean a local file path, pass pathlib.Path explicitly to avoid "
      "ambiguity with raw text."
  )
